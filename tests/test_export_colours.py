"""Export colours: order-number text by type, Glacier row fill, legend, TO flattening, truck grouping. Mocked, no Zoho."""
from __future__ import annotations

import os
from datetime import date
from types import SimpleNamespace

os.environ.setdefault("JWT_SECRET", "test")
os.environ.setdefault("ALLOWED_ORIGINS", "http://localhost")

from openpyxl import load_workbook  # noqa: E402

from services import inventory_exports as ex  # noqa: E402


def so(number="SO-1", lines=None, **extra):
    lines = lines if lines is not None else [{"item_id": "i1", "name": "Egg", "sku": "E1", "quantity": 2, "unit": "case", "warehouse_name": "Mets Cold Storage"}]
    return SimpleNamespace(salesorder_number=number, customer_name="Cust", shipping_address=None, expected_shipment_date=date(2026, 10, 14), order_status="confirmed", total=5,
                           raw_json={"line_items": lines, "shipping_address": {"address": "1 Road", "city": "Manila"}}, **extra)


def to(number="TO-1", source="mets"):
    name = ex.WAREHOUSES[source]["name"]
    return SimpleNamespace(salesorder_number=number, customer_name="Glacier Cold Storage", shipping_address={"address": "Amvel", "city": "Paranaque"}, expected_shipment_date=date(2026, 10, 14), order_status="draft", total=None, order_type="TO",
                           raw_json={"source_warehouse_key": source, "line_items": [{"item_id": "i2", "name": "Fish", "sku": "F1", "quantity": 4, "unit": "kg"}]})


def confirmed_rows(*orders, truck=("T1", "Juan")):
    rows = []
    for order in orders:
        rows += ex.flatten_confirmed_order(order, lambda _i: {"mets": 1, "glacier": 2}, lambda _o: truck)
    return rows


def order_cell(sheet, number):
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value == number:
                return cell
    raise AssertionError(number)


def test_order_number_font_colours_per_type_confirmed():
    rows = confirmed_rows(so("SO-1"), to("TO-1"))
    sheet = load_workbook(ex.make_excel(rows, "cap", "confirmed")).active
    assert order_cell(sheet, "SO-1").font.color.rgb.endswith("C00000")
    assert order_cell(sheet, "TO-1").font.color.rgb.endswith("2E7D32")
    # only the order-number cell is coloured: other cells in the row keep the default font colour
    row = order_cell(sheet, "SO-1").row
    assert sheet.cell(row, 2).font.color is None or not str(sheet.cell(row, 2).font.color.rgb).endswith("C00000")


def test_po_colour_is_blue():
    rows = confirmed_rows(so("PO-1", order_type="PO"))
    sheet = load_workbook(ex.make_excel(rows, "cap", "confirmed")).active
    assert order_cell(sheet, "PO-1").font.color.rgb.endswith("1F4E9E")


def test_type_column_present_in_both_layouts():
    assert ex.CONFIRMED_HEADERS[ex.CONFIRMED_TYPE_COLUMN] == "Type" and ex.CONFIRMED_HEADERS[ex.CONFIRMED_ORDER_COLUMN] == "Order Number"
    assert ex.HEADERS[ex.DEFAULT_TYPE_COLUMN] == "Type"
    assert len(ex.CONFIRMED_WIDTHS) == len(ex.CONFIRMED_HEADERS) and len(ex.DEFAULT_WIDTHS) == len(ex.HEADERS)
    crow = confirmed_rows(so())[0]
    drow = ex.flatten_order(so())[0]
    assert len(crow) == len(ex.CONFIRMED_HEADERS) and len(drow) == len(ex.HEADERS)
    assert crow[ex.CONFIRMED_TYPE_COLUMN] == "SO" and drow[ex.DEFAULT_TYPE_COLUMN] == "SO"  # missing order_type = SO


def test_default_layout_colours_and_glacier_fill():
    glacier = so("SO-G", [{"name": "Ice", "quantity": 1, "unit": "kg", "warehouse_name": "Glacier Cold Storage"}])
    rows = ex.flatten_order(so("SO-M")) + ex.flatten_order(glacier) + ex.flatten_order(to("TO-9", "glacier"))
    sheet = load_workbook(ex.make_excel(rows, "cap")).active
    assert order_cell(sheet, "TO-9").font.color.rgb.endswith("2E7D32")
    g = order_cell(sheet, "SO-G").row
    m = order_cell(sheet, "SO-M").row
    assert all(sheet.cell(g, c).fill.fgColor.rgb.endswith("DDEBF7") for c in range(1, len(ex.HEADERS) + 1))
    assert all(sheet.cell(m, c).fill.fill_type is None for c in range(1, len(ex.HEADERS) + 1))


def test_confirmed_glacier_fill_and_mets_no_fill_and_legend():
    glacier = so("SO-G", [{"name": "Ice", "quantity": 1, "unit": "kg", "warehouse_name": "Glacier Cold Storage"}])
    book = load_workbook(ex.make_excel(confirmed_rows(so("SO-M"), glacier), "cap", "confirmed"))
    sheet = book["Sales Orders"]
    g, m = order_cell(sheet, "SO-G").row, order_cell(sheet, "SO-M").row
    assert all(sheet.cell(g, c).fill.fgColor.rgb.endswith("DDEBF7") for c in range(1, len(ex.CONFIRMED_HEADERS) + 1))
    assert all(sheet.cell(m, c).fill.fill_type is None for c in range(1, len(ex.CONFIRMED_HEADERS) + 1))
    legend = " ".join(str(c.value) for row in book["Legend"].iter_rows() for c in row if c.value)
    assert "blue" in legend and "red" in legend and "green" in legend and "Glacier" in legend and "Mets" in legend


def test_warehouse_fill_can_be_disabled():
    glacier = so("SO-G", [{"name": "Ice", "quantity": 1, "unit": "kg", "warehouse_name": "Glacier Cold Storage"}])
    sheet = load_workbook(ex.make_excel(confirmed_rows(glacier), "cap", "confirmed", warehouse_fill=False))["Sales Orders"]
    assert sheet.cell(order_cell(sheet, "SO-G").row, 1).fill.fill_type is None


def test_to_flatten_uses_destination_and_source_warehouse():
    row = confirmed_rows(to("TO-5", "mets"))[0]
    assert row[ex.CONFIRMED_ORDER_COLUMN] == "TO-5" and row[ex.CONFIRMED_TYPE_COLUMN] == "TO"
    assert row[7] == "Glacier Cold Storage" and "Amvel" in row[ex.CONFIRMED_HEADERS.index("Shipping Address")]
    assert ex.warehouse_key(row[ex.CONFIRMED_WAREHOUSE_COLUMN]) == "mets"  # lines come from the source warehouse


def test_group_by_truck_then_type_then_number():
    rows = confirmed_rows(to("TO-1"), truck=("B", "")) + confirmed_rows(so("SO-2"), truck=("B", "")) + confirmed_rows(so("SO-1"), truck=("A", "")) + confirmed_rows(so("SO-0"), truck=("", ""))
    kinds = list(ex.group_by_truck(rows))
    headings = [p for k, p in kinds if k == "heading"]
    assert headings == ["A", "B", "No truck assigned"]
    numbers = [p[ex.CONFIRMED_ORDER_COLUMN] for k, p in kinds if k == "row"]
    assert numbers == ["SO-1", "SO-2", "TO-1", "SO-0"]
    sheet = load_workbook(ex.make_excel(rows, "cap", "confirmed"))["Sales Orders"]
    assert sheet.cell(2, 1).value == "Truck: A"


def test_pdf_builds_with_colour_markup():
    markup = ex.order_number_markup("TO-1", "TO")
    assert markup == '<font color="#2E7D32">TO-1</font>'
    assert ex.order_number_markup("<b>", "SO") == '<font color="#C00000">&lt;b&gt;</font>'
    assert 'color="#1F4E9E"' in ex.legend_markup()
    rows = confirmed_rows(so("SO-1"), to("TO-1"))
    for layout, data in (("confirmed", rows), ("default", ex.flatten_order(so("SO-1")) + ex.flatten_order(to("TO-1")))):
        pdf = ex.make_pdf(data, "cap", layout).getvalue()
        assert pdf.startswith(b"%PDF") and len(pdf) > 1000
