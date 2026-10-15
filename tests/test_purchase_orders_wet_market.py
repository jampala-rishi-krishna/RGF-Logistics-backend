"""Purchase orders (PO) and Wet Market (WM) orders in Load Planning: the PO- / SariSuki filter, statuses, expected-delivery date logic,
caching / Hard refresh, the WM split of the existing sales-order list (no extra Zoho call), assign / unassign, warehouse email rules,
export colours and storage. Mocked only: no Zoho, no database."""
from __future__ import annotations

import os
from datetime import date, datetime, timedelta
from types import SimpleNamespace
from zoneinfo import ZoneInfo

import pytest

os.environ.setdefault("JWT_SECRET", "test")
os.environ.setdefault("ALLOWED_ORIGINS", "http://localhost")

from openpyxl import load_workbook  # noqa: E402

from auth.dependencies import CurrentUser  # noqa: E402
from routers import assignment, load_planning  # noqa: E402
from services import (  # noqa: E402
    branches, confirmed_email, inventory_exports as ex, item_weight, live_sales_order_cache as lc, order_types, purchase_orders, sales_order_history_sync, transfer_orders,
)

RGF = "4489499000001322444"
SARISUKI = "4489499000044793937"
RARE_CUTS = "4489499000017304175"
METS = "4489499000064563124"
GLACIER = "4489499000064711027"
DAY = "2026-10-20"


def po(pid, number, *, status="open", delivery=DAY, day="2026-10-10", branch=RGF, branch_name="Rare Global Food Trading Corp.", location=GLACIER, vendor="Fresh Farms Inc.", total=1500):
    return {"purchaseorder_id": pid, "purchaseorder_number": number, "status": status, "date": day, "delivery_date": delivery, "branch_id": branch, "branch_name": branch_name,
            "location_id": location, "vendor_id": f"v{pid}", "vendor_name": vendor, "total": total, "last_modified_time": "2026-10-14T08:00:00+0800"}


RECORDS = [
    po("1", "PO-1026-01704"),
    po("2", "PO-1026-01703", status="billed", location=METS, vendor="Ocean Catch"),
    po("3", "SS PO-1026-12557"),  # "SS PO-": never
    po("4", "PO-1026-09999", branch=SARISUKI, branch_name="SariSuki Store Inc."),  # PO- but the SariSuki branch: never
    po("5", "PO-1026-01705", status="cancelled"),
    po("6", "PO-1026-01706", status="closed"),
    po("7", "PO-1026-01707", status="draft", delivery="", day=DAY),  # no expected date: planned on its PO date, flagged
    po("8", "PO-1026-01708", status="partially_received", delivery="2026-10-21"),
    po("9", "PO-1026-01709", branch=RARE_CUTS, branch_name="Rare Cuts"),  # not SariSuki: kept (the rule only drops SariSuki)
    po("10", "RC-PO-0001"),  # not a "PO-" number
]
DETAILS = {
    "1": {"line_items": [{"item_id": "i-kg", "name": "Pork belly", "sku": "PB", "quantity": 12, "unit": "kg", "warehouse_id": GLACIER}, {"item_id": "i-box", "name": "Chicken", "sku": "CH", "quantity": 4, "unit": "box"}],
          "billing_address": {"address": "12 Farm Road", "city": "Lipa", "state": "Batangas"}},
    "2": {"line_items": [{"item_id": "i-kg", "name": "Tuna", "sku": "TU", "quantity": 5, "unit": "kg"}], "billing_address": {"address": "Pier 4", "city": "Navotas"}},
}


class InlinePool:
    def submit(self, fn, *args, **kwargs):
        fn(*args, **kwargs)


class Zoho:
    def __init__(self):
        self.list_calls = 0
        self.detail_calls: list[str] = []


@pytest.fixture
def zoho(monkeypatch):
    purchase_orders.reset()
    fake = Zoho()

    def page(page=1, per_page=200):
        fake.list_calls += 1
        return {"purchaseorders": [dict(r) for r in RECORDS], "page_context": {"has_more_page": False}}

    def detail(pid):
        fake.detail_calls.append(pid)
        listed = next((dict(r) for r in RECORDS if r["purchaseorder_id"] == pid), {})
        return {"code": 0, "purchaseorder": {**listed, **DETAILS.get(pid, {"line_items": []}), "purchaseorder_id": pid}}

    monkeypatch.setattr(purchase_orders, "fetch_purchase_orders_page", page)
    monkeypatch.setattr(purchase_orders, "fetch_purchase_order_detail", detail)
    monkeypatch.setattr(purchase_orders, "_pool", InlinePool())
    monkeypatch.setitem(item_weight._cache, "i-box", 2.5)
    yield fake
    purchase_orders.reset()
    with lc._state_lock:
        for oid in [o for o in lc._assignment_state if str(o).startswith(("PO:", "so-", "wm-"))]:
            lc._assignment_state.pop(oid, None)
    lc._assigned_zoho.clear()


# ---- the order-type rule ----------------------------------------------------------------------------------------------------


def test_order_type_of_every_kind():
    assert order_types.order_type_of("PO:1") == "PO" and order_types.order_type_of("TO:1") == "TO"
    assert order_types.order_type_of("123", "WM-SO26-00282") == "WM" and order_types.order_type_of("123", "SO-1") == "SO"
    assert order_types.order_type_of(SimpleNamespace(id="5", salesorder_number="WM-SO26-00282", order_type="SO")) == "WM"  # saved before the split
    assert order_types.order_type_of(SimpleNamespace(id="5", salesorder_number="SS SO26-16612")) == "SO"
    assert order_types.order_type_of({"id": "x", "order_type": "PO"}) == "PO"
    assert order_types.parse_types(None) == {"SO", "WM"} and order_types.parse_types("so,po,xx") == {"SO", "PO"}


# ---- PO filter --------------------------------------------------------------------------------------------------------------


@pytest.mark.parametrize("number,branch,branch_name,wanted", [
    ("PO-1026-01704", RGF, "Rare Global Food Trading Corp.", True),
    ("po-1026-01704", RGF, "", True),
    ("SS PO-1026-12557", RGF, "", False),
    ("SS PO-1026-12557", SARISUKI, "SariSuki Store Inc.", False),
    ("PO-1026-01704", SARISUKI, "SariSuki Store Inc.", False),
    ("PO-1026-01704", "", "SariSuki Store Inc.", False),  # by branch name when the id is missing
    ("RC-PO-1", RGF, "", False),
])
def test_po_prefix_and_branch_filter(number, branch, branch_name, wanted):
    assert purchase_orders.is_wanted({"purchaseorder_number": number, "branch_id": branch, "branch_name": branch_name}) is wanted


def test_pull_keeps_only_po_numbers_outside_sarisuki_and_counts_statuses(zoho):
    records, _ = purchase_orders.get_records()
    assert sorted(records, key=int) == ["1", "2", "5", "6", "7", "8", "9"]
    assert purchase_orders.loaded_status_counts() == {"open": 2, "billed": 1, "cancelled": 1, "closed": 1, "draft": 1, "partially_received": 1}


def test_default_hides_cancelled_and_closed_and_the_status_filter_shows_them(zoho):
    rows, _, _, counts = purchase_orders.list_rows(date(2026, 10, 20), date(2026, 10, 20))
    assert sorted(r.salesorder_number for r in rows) == ["PO-1026-01703", "PO-1026-01704", "PO-1026-01707", "PO-1026-01709"]
    assert counts == {"open": 2, "billed": 1, "cancelled": 1, "closed": 1, "draft": 1}  # before the status filter, for the tab's options
    rows, *_ = purchase_orders.list_rows(date(2026, 10, 20), date(2026, 10, 20), statuses=purchase_orders.parse_statuses("Cancelled,closed"))
    assert sorted(r.salesorder_number for r in rows) == ["PO-1026-01705", "PO-1026-01706"]
    rows, *_ = purchase_orders.list_rows(date(2026, 10, 20), date(2026, 10, 20), statuses=purchase_orders.parse_statuses("all"))
    assert len(rows) == 6


def test_expected_delivery_date_is_the_date_and_falls_back_to_the_po_date(zoho):
    rows, *_ = purchase_orders.list_rows(date(2026, 10, 21), date(2026, 10, 21))
    assert [r.salesorder_number for r in rows] == ["PO-1026-01708"]  # PO date 10-10, expected delivery 10-21
    assert purchase_orders.list_rows(date(2026, 10, 10), date(2026, 10, 10))[0] == []  # never planned on its PO date when it has an expected date
    rows, *_ = purchase_orders.list_rows(date(2026, 10, 20), date(2026, 10, 20))
    flagged = {r.salesorder_number: purchase_orders.summary(r)["no_expected_date"] for r in rows}
    assert flagged["PO-1026-01707"] is True and flagged["PO-1026-01704"] is False


def test_list_is_cached_date_change_costs_nothing_hard_refresh_bypasses(zoho):
    purchase_orders.list_rows(date(2026, 10, 20), date(2026, 10, 20))
    purchase_orders.list_rows(date(2026, 10, 21), date(2026, 10, 21))
    assert zoho.list_calls == 1
    purchase_orders.get_records(force=True)
    assert zoho.list_calls == 2


def test_details_only_for_shown_orders_and_kg(zoho):
    purchase_orders.list_rows(date(2026, 10, 20), date(2026, 10, 20))
    assert sorted(zoho.detail_calls, key=int) == ["1", "2", "7", "9"]  # cancelled / closed / other days: no detail call
    rows, _, pending, _ = purchase_orders.list_rows(date(2026, 10, 20), date(2026, 10, 20))
    assert pending == 0 and len(zoho.detail_calls) == 4  # nothing fetched twice
    first = purchase_orders.summary(next(r for r in rows if r.id == "PO:1"))
    assert [p["total_weight_kg"] for p in first["products"]] == [12.0, 10.0]
    assert purchase_orders.total_weight_kg([r for r in rows if r.id in ("PO:1", "PO:2")]) == (27.0, True, [])


def test_kg_plus_while_lines_are_unknown(zoho, monkeypatch):
    monkeypatch.setattr(purchase_orders, "request_details", lambda records: len(records))
    rows, *_ = purchase_orders.list_rows(date(2026, 10, 20), date(2026, 10, 20))
    assert purchase_orders.total_weight_kg(rows)[:2] == (0.0, False)


def test_po_row_shape_vendor_pickup_and_destination_warehouse(zoho):
    purchase_orders.list_rows(date(2026, 10, 20), date(2026, 10, 20))
    rows, *_ = purchase_orders.list_rows(date(2026, 10, 20), date(2026, 10, 20))
    by_id = {r.id: purchase_orders.summary(r) for r in rows}
    one, two = by_id["PO:1"], by_id["PO:2"]
    assert one["order_type"] == "PO" and one["purchaseorder_number"] == "PO-1026-01704" and one["vendor_name"] == "Fresh Farms Inc."
    assert (one["destination_warehouse"], one["destination_warehouse_key"]) == ("Glacier Cold Storage", "glacier")  # by warehouse ID
    assert (two["destination_warehouse"], two["destination_warehouse_key"]) == ("Mets Cold Storage", "mets")
    assert one["pickup_from"] == "Fresh Farms Inc." and one["shipping_address"]["city"] == "Lipa"  # the stop = the vendor
    assert one["total"] == 1500.0 and one["status_label"] == "Open" and one["branch_code"] == "RGF" and one["expected_delivery_date"] == DAY
    assert one["zoho_lock"] is None and one["mets_qty_available_for_sale"] is None


def test_destination_falls_back_to_the_warehouse_name():
    assert purchase_orders.destination({"location_id": "unknown", "location_name": "Mets Cold Storage - RGF"}) == ("mets", "Mets Cold Storage - RGF")
    assert purchase_orders.destination({"location_name": "Production Area RC"}) == (None, "Production Area RC")


# ---- PO endpoints ----------------------------------------------------------------------------------------------------------


def call_po(**overrides):
    args = dict(date_from=DAY, date_to=DAY, search=None, statuses=None, assignment=None, page=1, per_page=25, db=None)
    args.update(overrides)
    return load_planning.list_purchase_orders(**args)


def test_po_endpoint_shape_status_counts_search_and_refresh(zoho):
    call_po()
    body = call_po()
    assert body["total"] == 4 and body["source"] == "zoho" and body["last_refreshed_at"] and body["hidden_by_default"] == ["cancelled", "closed"]
    assert body["status_counts"]["cancelled"] == 1 and body["status_counts"]["open"] == 2
    assert {i["order_type"] for i in body["items"]} == {"PO"}
    assert call_po(search="ocean")["items"][0]["id"] == "PO:2"
    assert call_po(search="po102601704")["items"][0]["id"] == "PO:1"
    assert call_po(statuses="closed")["total"] == 1
    before = zoho.list_calls
    out = load_planning.hard_refresh_purchase_orders()
    assert zoho.list_calls == before + 1 and out["count"] == 7 and out["status_counts"]["billed"] == 1


# ---- WM: the same sales-order list, split by number ------------------------------------------------------------------------


def _today() -> date:
    return datetime.now(ZoneInfo("Asia/Manila")).date()


def so_record(rid, number, branch=RGF):
    day = (_today() + timedelta(days=1)).isoformat()
    return {"salesorder_id": rid, "salesorder_number": number, "customer_name": f"Cust {rid}", "status": "confirmed", "shipment_date": day, "date": day, "branch_id": branch,
            "last_modified_time": "2026-10-14T08:00:00+0800", "line_items": [{"item_id": "i-kg", "name": "Fish", "quantity": 3, "unit": "kg"}]}


@pytest.fixture
def window(monkeypatch):
    calls = {"list": 0, "detail": 0}
    records = [so_record("101", "SO-00101"), so_record("102", "WM-SO26-00282"), so_record("103", "WM-SO26-00283"), so_record("104", "MS-SO-00104")]

    def by_date(start, end, page=1, per_page=200, sort_column=None):
        calls["list"] += 1
        return {"salesorders": [dict(r) for r in records], "page_context": {"has_more_page": False, "search_criteria": [{"column_name": "shipment_date"}]}}

    def detail(rid):
        calls["detail"] += 1
        return {"salesorder": next(dict(r) for r in records if r["salesorder_id"] == rid)}

    monkeypatch.setattr(lc, "fetch_sales_orders_by_shipment_date", by_date)
    monkeypatch.setattr(lc, "fetch_sales_order_detail", detail)
    monkeypatch.setattr(lc, "FORCE_BLOCKING_READS", True)
    monkeypatch.setattr(load_planning, "_ack_snapshot", lambda: (set(), {}))
    monkeypatch.setattr(load_planning, "stock_for_orders_cached", lambda rows: ({}, 0))
    monkeypatch.setattr(load_planning, "cached_item_stock", lambda *a, **k: None)
    with lc._window_lock:
        lc._windows.clear()
    yield calls
    with lc._window_lock:
        lc._windows.clear()
    lc._details.clear()


def call_so(order_types_value):
    day = (_today() + timedelta(days=1)).isoformat()
    return load_planning.list_sales_orders(date_from=day, date_to=day, status="All except acknowledged", search=None, assignment=None, cities=None, vehicle=None, customer=None,
                                           delivery_status=None, branches=None, order_types=order_types_value, page=1, per_page=100, db=None)


def test_wm_split_so_tab_excludes_wm_and_wm_tab_has_only_wm_with_no_extra_zoho_calls(window):
    so_tab = call_so("SO")
    after_so = dict(window)
    wm_tab = call_so("WM")
    assert sorted(i["salesorder_number"] for i in so_tab["items"]) == ["MS-SO-00104", "SO-00101"]
    assert sorted(i["salesorder_number"] for i in wm_tab["items"]) == ["WM-SO26-00282", "WM-SO26-00283"]
    assert {i["order_type"] for i in wm_tab["items"]} == {"WM"} and {i["order_type"] for i in so_tab["items"]} == {"SO"}
    assert window == after_so and window["list"] == 1  # the WM tab is served from the SO tab's list: zero extra Zoho calls
    assert call_so(None)["total"] == 4  # absent filter: every sales order (SO + WM), as before


def test_ss_so_is_judged_by_branch_not_by_number():
    # SS SO numbers are SariSuki's series: they are excluded by the SariSuki branch id (the allowlist), not by the number. An "SS SO"
    # carried on an allowed branch would show in the SO tab - nothing here filters it by prefix.
    assert branches.is_allowed({"salesorder_number": "SS SO26-16612", "branch_id": SARISUKI}) is False
    assert branches.is_allowed({"salesorder_number": "SS SO26-16612", "branch_id": RGF}) is True
    assert order_types.order_type_of("1", "SS SO26-16612") == "SO"


def test_saved_po_and_wm_rows_pass_the_branch_allowlist():
    assert branches.is_allowed(SimpleNamespace(id="PO:1", salesorder_number="PO-1026-01704", order_type="PO", raw_json=None))
    assert branches.is_allowed(SimpleNamespace(id="7", salesorder_number="WM-SO26-00282", order_type="WM", raw_json=None))


# ---- assign / unassign ------------------------------------------------------------------------------------------------------


class Flow:
    def __init__(self, monkeypatch):
        self.user = CurrentUser(id=1, email="pau@example.com", role="dispatcher", full_name="Pau", status="active")
        profile = SimpleNamespace(plate_no="TRK1", rated_capacity_kg=1000, is_reefer=None, is_third_party=False, capacity_note=None)
        self.added = []
        self.db = SimpleNamespace(execute=lambda stmt: SimpleNamespace(scalar_one_or_none=lambda: profile, scalars=lambda: SimpleNamespace(all=lambda: []), all=lambda: []),
                                  commit=lambda: None, add=self.added.append)
        self.wm = SimpleNamespace(id="wm-1", salesorder_number="WM-SO26-00282", customer_name="Market stall", shipping_address=None, order_status="confirmed", manifest_id=None, route_id=None,
                                  raw_json={"shipping_address": {"city": "Manila"}, "line_items": [{"quantity": 7, "unit": "kg"}]}, assignment_status="unassigned",
                                  expected_shipment_date=date(2026, 10, 20), vehicle_id=None, driver_id=None, helper_ids=[], assigned_at=None, assigned_by=None, order_type="WM")
        lc._assigned_zoho["wm-1"] = self.wm
        monkeypatch.setattr(assignment.staff_directory_cache, "get_by_id", lambda i: {"id": i, "name": "Juan"})
        monkeypatch.setattr(assignment, "sync_history_row", lambda db, o: None)
        monkeypatch.setattr(assignment, "invalidate_fleet_cache", lambda: None)
        monkeypatch.setattr(lc, "get_assigned_snapshot", lambda: list(lc._assigned_zoho.values()))
        monkeypatch.setattr(lc, "invalidate_assigned_zoho", lambda: None)
        monkeypatch.setattr(lc, "invalidate_windows", lambda: None)

    def assign(self, ids):
        return assignment.assign_order(ids[0], assignment.AssignmentBody(salesorder_ids=ids, vehicle_id="TRK1", driver_ids=[1]), self.user, self.db)


@pytest.fixture
def flow(zoho, monkeypatch):
    return Flow(monkeypatch)


def test_assign_po_and_wm_together_leaves_the_po_tab_and_unassign_brings_it_back(zoho, flow):
    assert call_po()["total"] == 4
    result = flow.assign(["PO:1", "wm-1"])
    assert result["success"] is True and sorted(lc.ids_for_batch(result["assignment_batch_id"])) == ["PO:1", "wm-1"]
    assert sorted(i["id"] for i in call_po()["items"]) == ["PO:2", "PO:7", "PO:9"]
    assert flow.wm.assignment_status == "assigned"
    assignment.unassign_order("PO:1", flow.user, flow.db, assignment.UnassignBody(confirm=True))
    assert "PO:1" in [i["id"] for i in call_po()["items"]]


def test_cannot_assign_a_po_outside_the_rule(zoho, flow):
    with pytest.raises(Exception) as caught:
        flow.assign(["PO:4"])  # SariSuki
    assert getattr(caught.value, "status_code", None) in (404, 409)
    assert lc.get_assignment("PO:4") is None


def test_bulk_unassign_counts_po_and_wm(zoho, flow):
    flow.assign(["PO:1", "PO:2", "wm-1"])
    out = assignment.bulk_unassign(assignment.BulkUnassignBody(order_ids=["PO:1", "wm-1"], confirm=True), flow.user, flow.db)
    assert out["count"] == 2 and out["count_by_type"] == {"SO": 0, "WM": 1, "PO": 1, "TO": 0}
    assert flow.wm.assignment_status == "unassigned"


def test_confirmed_orders_lists_all_four_types(zoho, flow, monkeypatch):
    flow.assign(["PO:1", "wm-1"])
    monkeypatch.setattr(load_planning, "_ack_snapshot", lambda: (set(), {}))
    monkeypatch.setattr(lc, "get_assigned_snapshot_ex", lambda block=True, **kw: (list(lc._assigned_zoho.values()), False))
    future = (_today() + timedelta(days=30)).isoformat()
    for row in lc._assigned_zoho.values():
        row.expected_shipment_date = date.fromisoformat(future)

    def rows(types):
        return sorted(r.id for r in load_planning._rows_any_branch(None, future, future, None, None, "assigned", order_types=types))

    assert rows("SO,WM,PO,TO") == ["PO:1", "wm-1"]
    assert rows("WM") == ["wm-1"] and rows("PO") == ["PO:1"] and rows("SO") == []
    assert rows(None) == ["wm-1"]  # absent: sales orders (SO + WM)


# ---- storage, driver, email, export ----------------------------------------------------------------------------------------


class FakeDb:
    def __init__(self):
        self.added = []

    def get(self, model, key):
        return None

    def add(self, obj):
        self.added.append(obj)

    def execute(self, stmt):
        return None


def test_history_row_stores_wm_and_po_order_type(zoho, monkeypatch):
    monkeypatch.setattr(sales_order_history_sync, "stock_for_orders", lambda orders: {})
    rows, *_ = purchase_orders.list_rows(date(2026, 10, 20), date(2026, 10, 20))
    rows, *_ = purchase_orders.list_rows(date(2026, 10, 20), date(2026, 10, 20))
    po_row = next(r for r in rows if r.id == "PO:1")
    db = FakeDb()
    sales_order_history_sync.sync_history_row(db, po_row)
    saved = db.added[0]
    assert (saved.order_type, saved.source_warehouse, saved.destination_warehouse, saved.customer_name) == ("PO", "Fresh Farms Inc.", "Glacier Cold Storage", "Fresh Farms Inc.")
    wm_row = lc._build_transient(so_record("102", "WM-SO26-00282"))
    db = FakeDb()
    sales_order_history_sync.sync_history_row(db, wm_row)
    assert db.added[0].order_type == "WM" and db.added[0].source_warehouse is None


def test_po_checkout_uses_the_po_number_as_reference():
    from services import trip_workflow

    row = SimpleNamespace(id="PO:1", salesorder_number="PO-1026-01704", order_type="PO")
    assert trip_workflow._invoice_no(row, None) == "PO-1026-01704"


def _order(kind, number, raw, lines):
    return SimpleNamespace(id=number, salesorder_number=number, customer_name="X", shipping_address=None, expected_shipment_date=date(2026, 10, 20), order_status="open", total=1,
                           order_type=kind, vehicle_id="TRK1", driver_id=None, helper_ids=[], raw_json={**raw, "line_items": lines})


def test_email_warehouse_rules_po_destination_to_source_so_wm_line():
    po_glacier = _order("PO", "PO-1", {"destination_warehouse_key": "glacier", "to_warehouse_name": "Glacier Cold Storage"}, [{"name": "Tuna", "quantity": 1, "unit": "kg", "warehouse_name": "Mets Cold Storage"}])
    to_mets = _order("TO", "TO-1", {"source_warehouse_key": "mets", "from_warehouse_name": "Mets Cold Storage"}, [{"name": "Pork", "quantity": 1, "unit": "kg"}])
    wm_glacier = _order("WM", "WM-SO26-1", {}, [{"name": "Crab", "quantity": 1, "unit": "kg", "warehouse_name": "Glacier Cold Storage"}])
    so_mixed = _order("SO", "SO-1", {}, [{"name": "A", "quantity": 1, "unit": "kg", "warehouse_name": "Mets Cold Storage"}, {"name": "B", "quantity": 1, "unit": "kg", "warehouse_name": "Glacier Cold Storage"}])
    everything = [po_glacier, to_mets, wm_glacier, so_mixed]
    glacier = {o.salesorder_number: len(o.raw_json["line_items"]) for o in confirmed_email.filter_orders_for_warehouse(everything, "glacier")}
    mets = {o.salesorder_number: len(o.raw_json["line_items"]) for o in confirmed_email.filter_orders_for_warehouse(everything, "mets")}
    assert glacier == {"PO-1": 1, "WM-SO26-1": 1, "SO-1": 1}  # the PO goes with its DESTINATION, whatever its line says
    assert mets == {"TO-1": 1, "SO-1": 1}  # the TO goes with its SOURCE


def _cell(sheet, value):
    return next(c for row in sheet.iter_rows() for c in row if c.value == value)


def test_export_colours_wm_purple_po_blue_so_red_to_green_and_legend():
    lines = [{"item_id": "i1", "name": "Fish", "sku": "F", "quantity": 1, "unit": "kg", "warehouse_name": "Glacier Cold Storage"}]
    orders = [_order("SO", "SO-1", {}, lines), _order("WM", "WM-SO26-00282", {}, lines), _order("PO", "PO-1026-01704", {"to_warehouse_name": "Glacier Cold Storage"}, lines),
              _order("TO", "TO-1", {}, lines), SimpleNamespace(**{**vars(_order("SO", "WM-SO26-00283", {}, lines)), "order_type": "SO"})]  # saved as SO before the split
    rows = [r for o in orders for r in ex.flatten_confirmed_order(o, lambda _i: {}, lambda _o: ("TRK1", "Juan"))]
    for layout, rows_ in (("confirmed", rows), ("default", [r for o in orders for r in ex.flatten_order(o)])):
        sheet = load_workbook(ex.make_excel(rows_, "cap", layout)).active
        assert _cell(sheet, "SO-1").font.color.rgb.endswith("C00000")
        assert _cell(sheet, "WM-SO26-00282").font.color.rgb.endswith("7C3AED")
        assert _cell(sheet, "WM-SO26-00283").font.color.rgb.endswith("7C3AED")
        assert _cell(sheet, "PO-1026-01704").font.color.rgb.endswith("1F4E9E")
        assert _cell(sheet, "TO-1").font.color.rgb.endswith("2E7D32")
        fill = _cell(sheet, "PO-1026-01704").fill.fgColor.rgb
        assert str(fill).endswith(ex.GLACIER_FILL)  # Glacier rows light blue, every type
    book = load_workbook(ex.make_excel(rows, "cap", "confirmed"))
    legend = " ".join(str(c.value) for row in book["Legend"].iter_rows() for c in row if c.value)
    assert "purple" in legend and "Wet Market" in legend
    assert 'color="#7C3AED"' in ex.legend_markup() and "WM (purple)" in ex.legend_markup()
    assert ex.order_number_markup("WM-SO26-1", "WM") == '<font color="#7C3AED">WM-SO26-1</font>'
    assert ex.make_pdf(rows, "cap", "confirmed").getvalue().startswith(b"%PDF")
