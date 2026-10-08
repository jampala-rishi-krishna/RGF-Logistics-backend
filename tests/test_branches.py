"""Branch allowlist: RareChain handles ONLY RGF and Meat and Seafood Specialist Inc. (MSSI).

Orders of any other branch (SariSuki, Rare Cuts, Rare Food Shop, never-seen branches) are never loaded,
counted, cached, hydrated, listed, exported or acted on. Zoho is faked in memory; the fake deliberately
IGNORES the branch_ids parameter, so these tests also prove the local drop works on its own.
"""
import os
import unittest
from types import SimpleNamespace
from unittest.mock import patch

os.environ.setdefault("JWT_SECRET", "test")

from routers import assignment, load_planning
from services import branches as bs
from services import live_sales_order_cache as lc
from services import warehouse_stock as ws
from services import zoho_client
from tests.test_inventory_ack_hold_capacity import SHIP, ZohoBackedTestCase

RGF, MSSI, SSI, RC, RFS = (b["id"] for b in bs.BRANCHES)
UNKNOWN = "999000111"
NOT_HANDLED = "This branch isn't handled in RareChain"

ALLOWED_NUMBERS = ["MS-SO-02016", "MS-SO-02017", "MS-SO-02099", "SO26-18001", "SO26-18002", "SO26-99001", "WM-SO26-00247"]
EXCLUDED_NUMBERS = {"SS SO26-16612", "SS SO26-16613", "RC-0001", "RFS-0001", "XX-1", "ZZ-5"}


def order(oid, number, branch_id, branch_name, sub="confirmed"):
    record = {"salesorder_id": oid, "salesorder_number": number, "customer_name": f"Cust {oid}", "status": "confirmed", "current_sub_status": sub, "order_sub_status": sub,
              "shipment_date": SHIP, "last_modified_time": "t0", "line_items": []}
    if branch_id:
        record.update(branch_id=branch_id, branch_name=branch_name)
    return record


class BranchTestCase(ZohoBackedTestCase):
    COUNT = 0

    def setUp(self):
        super().setUp()
        self.detail_ids = []
        self.zoho.orders = {o["salesorder_id"]: o for o in (
            order("r1", "SO26-18001", RGF, "Rare Global Food Trading Corp."),
            order("r2", "SO26-18002", RGF, "Rare Global Food Trading Corp."),
            order("w1", "WM-SO26-00247", RGF, "Rare Global Food Trading Corp."),
            order("m1", "MS-SO-02016", MSSI, "Meat and Seafood Specialist Inc."),
            order("m2", "MS-SO-02017", MSSI, "Meat and Seafood Specialist Inc."),
            order("s1", "SS SO26-16612", SSI, "SariSuki Store Inc."),
            order("s2", "SS SO26-16613", SSI, "SariSuki Store Inc.", sub="cs_acknowl"),
            order("c1", "RC-0001", RC, "Rare Cuts"),
            order("f1", "RFS-0001", RFS, "Rare Food Shop"),
            order("u1", "XX-1", UNKNOWN, "Foo Store"),
            order("n1", "SO26-99001", None, None),       # no branch id, RGF numbered -> allowed
            order("n2", "ZZ-5", None, None),             # no branch id, not RGF/MSSI numbered -> excluded
            order("n3", "MS-SO-02099", None, None),      # no branch id, MSSI numbered -> allowed
        )}
        detail = self.zoho.detail

        def recording_detail(order_id):
            self.detail_ids.append(order_id)
            return detail(order_id)

        # Zoho's Acknowledged custom view covers RGF + MSSI only.
        self.zoho.view = lambda customview_id, page=1, per_page=200: {"salesorders": [{"salesorder_id": i, "branch_id": o.get("branch_id"), "salesorder_number": o["salesorder_number"]} for i, o in sorted(self.zoho.orders.items()) if o["current_sub_status"] == "cs_acknowl" and o.get("branch_id") in (RGF, MSSI)], "page_context": {"has_more_page": False}}
        for target, attr, fake in ((load_planning, "fetch_sales_orders_by_customview", self.zoho.view), (lc, "fetch_sales_order_detail", recording_detail), (load_planning, "fetch_sales_order_detail", recording_detail)):
            patcher = patch.object(target, attr, fake)
            patcher.start()
            self.addCleanup(patcher.stop)
        lc._excluded_ids.clear()
        self.addCleanup(lc._excluded_ids.clear)

    def listing(self, branches=None, status="All except acknowledged", assignment=None, search=None):
        return load_planning.list_sales_orders(date_from=SHIP, date_to=SHIP, status=status, search=search, assignment=assignment, cities=None, vehicle=None, customer=None, delivery_status=None, branches=branches, page=1, per_page=100, db=SimpleNamespace())


class ExcludedBranchesNeverAppearTests(BranchTestCase):
    def test_inventory_lists_only_rgf_and_mssi(self):
        page = self.listing()
        self.assertEqual(self.numbers(page), sorted(ALLOWED_NUMBERS))
        self.assertEqual(page["total"], len(ALLOWED_NUMBERS))

    def test_ms_orders_still_visible_with_badge(self):
        items = {i["salesorder_number"]: i for i in self.listing()["items"]}
        for number in ("MS-SO-02016", "MS-SO-02017"):
            self.assertEqual((items[number]["branch_code"], items[number]["branch_id"]), ("MSSI", MSSI))
        self.assertEqual(items["WM-SO26-00247"]["branch_code"], "RGF")  # WM-SO is an RGF series

    def test_unknown_branch_and_unnumbered_unknowns_are_excluded(self):
        self.assertFalse(EXCLUDED_NUMBERS & set(self.numbers(self.listing())))

    def test_orders_without_branch_id_follow_their_number(self):
        numbers = self.numbers(self.listing())
        self.assertIn("SO26-99001", numbers)   # RGF numbered
        self.assertIn("MS-SO-02099", numbers)  # MSSI numbered
        self.assertNotIn("ZZ-5", numbers)

    def test_counts_only_cover_allowed_branches(self):
        self.assertEqual(self.listing()["branch_counts"], {RGF: 3, MSSI: 2})  # no-branch rows are not attributed
        self.assertEqual(self.listing(branches=f"{MSSI},{SSI}")["branch_counts"], {RGF: 3, MSSI: 2})

    def test_asking_for_an_excluded_branch_returns_nothing(self):
        for excluded in (SSI, RC, RFS, UNKNOWN):
            self.assertEqual(self.listing(branches=excluded)["total"], 0)

    def test_search_cannot_surface_excluded_orders(self):
        for text in ("SS SO26-16612", "ss-so26-16612", "SS SO26", "RC-0001", "XX-1", "ZZ-5"):
            self.assertEqual(self.listing(search=text)["total"], 0, text)

    def test_no_detail_gets_for_excluded_orders(self):
        self.listing()
        self.assertTrue(self.detail_ids)
        self.assertFalse({"s1", "s2", "c1", "f1", "u1", "n2"} & set(self.detail_ids))

    def test_load_planning_and_acknowledged_views_exclude_them(self):
        lp = load_planning._filtered_rows(SimpleNamespace(), SHIP, SHIP, None, None, "unassigned")
        self.assertFalse(EXCLUDED_NUMBERS & {r.salesorder_number for r in lp})
        # the only acknowledged order is the SSI one, which is excluded
        self.assertEqual(self.listing(status="Acknowledged", assignment="unassigned")["total"], 0)

    def test_export_email_and_bulk_acknowledge_exclude_them(self):
        rows = load_planning._filtered_rows(SimpleNamespace(), SHIP, SHIP, "All except acknowledged", None, None)
        self.assertEqual(sorted(r.salesorder_number for r in rows), sorted(ALLOWED_NUMBERS))
        ctx = load_planning.EmailFilterContext(date_from=SHIP, date_to=SHIP, status="All except acknowledged")
        with patch.object(load_planning, "_hydrate_export_rows", lambda db, rows: rows):
            self.assertEqual(sorted(r.salesorder_number for r in load_planning._email_rows(SimpleNamespace(), ctx)), sorted(ALLOWED_NUMBERS))
        result = load_planning.acknowledge_filtered_sales_orders(date_from=SHIP, date_to=SHIP, status="All except acknowledged", search=None, branches=None, db=SimpleNamespace(), current_user=self.user)
        self.assertEqual(result["eligible_count"], len(ALLOWED_NUMBERS))
        self.assertEqual(self.zoho.acked() & {"s1", "c1", "f1", "u1", "n2"}, set())

    def test_assigned_snapshot_hides_excluded_without_counting_a_failure(self):
        lc._assignment_state.clear()
        lc._assigned_zoho.clear()
        self.addCleanup(lc._assignment_state.clear)
        self.addCleanup(lc._assigned_zoho.clear)
        for oid in ("r1", "m1", "s1"):
            lc._assignment_state[oid] = {"assignment_status": "assigned", "vehicle_id": "V1"}
        rows, had_failures = lc.get_assigned_snapshot_ex()
        self.assertEqual(sorted(r.salesorder_number for r in rows), ["MS-SO-02016", "SO26-18001"])
        self.assertFalse(had_failures)
        self.assertIn("s1", lc._excluded_ids)

    def test_past_dated_history_hides_excluded_numbers(self):
        rows = [SimpleNamespace(id=n, salesorder_number=n, raw_json=None, assignment_status="assigned", expected_shipment_date=None) for n in ("SO26-1", "WM-SO26-2", "SS SO26-3", "ZZ-4")]
        db = SimpleNamespace(execute=lambda query: SimpleNamespace(scalars=lambda: SimpleNamespace(all=lambda: rows)))
        got = load_planning._filtered_rows(db, "2020-01-01", "2020-01-02", None, None, "assigned")
        self.assertEqual([r.salesorder_number for r in got], ["SO26-1", "WM-SO26-2"])


class BranchOptionsTests(BranchTestCase):
    def test_options_are_rgf_and_mssi_only(self):
        self.assertEqual([b["id"] for b in load_planning.list_branches()["branches"]], [RGF, MSSI])
        self.assertEqual([b["code"] for b in self.listing()["branches"]], ["RGF", "MSSI"])
        # an unknown branch seen in data is never offered
        self.assertEqual([b["id"] for b in bs.all_options([{"id": UNKNOWN, "name": "Foo", "code": "F", "label": "Foo"}])], [RGF, MSSI])

    def test_filter_by_rgf_or_mssi_still_works(self):
        self.assertEqual(self.numbers(self.listing(branches=MSSI)), ["MS-SO-02016", "MS-SO-02017"])
        self.assertEqual(self.numbers(self.listing(branches=RGF)), ["SO26-18001", "SO26-18002", "WM-SO26-00247"])
        self.assertEqual(self.listing(branches=f"{RGF},{MSSI}")["total"], 5)

    def test_allowlist_is_overridable_by_env(self):
        with patch.object(bs, "ALLOWED_BRANCH_IDS", (RGF,)):
            self.assertEqual(self.numbers(self.listing()), ["SO26-18001", "SO26-18002", "SO26-99001", "WM-SO26-00247"])
        with patch.dict(os.environ, {"ALLOWED_BRANCHES": f"{MSSI}, {SSI}"}):
            self.assertEqual(bs._resolve_allowed(), (MSSI, SSI))
        with patch.dict(os.environ, {"ALLOWED_BRANCHES": ""}):
            self.assertEqual(bs._resolve_allowed(), (RGF, MSSI))
        self.assertEqual(bs.ALLOWED_BRANCH_IDS, (RGF, MSSI))


class NumberSearchTests(BranchTestCase):
    def found(self, text):
        return self.numbers(self.listing(search=text))

    def test_every_handled_prefix_format(self):
        self.assertEqual(self.found("MS-SO-02016"), ["MS-SO-02016"])
        self.assertEqual(self.found("WM-SO26-00247"), ["WM-SO26-00247"])
        self.assertEqual(self.found("SO26-18001"), ["SO26-18001"])

    def test_punctuation_and_case_insensitive(self):
        self.assertEqual(self.found("ms so 02017"), ["MS-SO-02017"])
        self.assertEqual(self.found("wmso2600247"), ["WM-SO26-00247"])


class WriteGuardTests(BranchTestCase):
    def assert_blocked(self, call):
        with self.assertRaises(bs.BranchNotAllowed) as caught:
            call()
        self.assertEqual((caught.exception.status_code, caught.exception.detail), (409, NOT_HANDLED))

    def test_acknowledge_is_rejected_for_every_excluded_branch(self):
        for oid in ("s1", "s2", "c1", "f1", "u1"):
            with self.subTest(order=oid):
                self.assert_blocked(lambda oid=oid: load_planning.acknowledge_sales_order_route(oid, self.user))
        self.assertEqual(self.zoho.calls["ack_post"], 0)

    def test_remove_acknowledge_lock_and_drawer_are_rejected(self):
        self.assert_blocked(lambda: load_planning.remove_acknowledge_sales_order_route("s2"))
        self.assert_blocked(lambda: load_planning.lock_sales_order_route("s2", self.user))
        self.assert_blocked(lambda: load_planning.get_sales_order("s1"))
        self.assertEqual(self.zoho.calls["remove_post"], 0)

    def test_known_excluded_order_is_rejected_without_another_zoho_call(self):
        self.assert_blocked(lambda: load_planning.acknowledge_sales_order_route("s1", self.user))
        before = len(self.detail_ids)
        self.assert_blocked(lambda: load_planning.acknowledge_sales_order_route("s1", self.user))
        self.assert_blocked(lambda: load_planning.get_sales_order("s1"))
        self.assertEqual(len(self.detail_ids), before)

    def test_assign_and_unassign_are_rejected(self):
        body = assignment.AssignmentBody(salesorder_ids=["s1"], vehicle_id="V1")
        self.assert_blocked(lambda: assignment.assign_order("s1", body, self.user, SimpleNamespace()))
        self.assert_blocked(lambda: assignment.unassign_order("s1", self.user, SimpleNamespace()))

    def test_handled_branches_still_work(self):
        out = self.acknowledge("m1", "r1")
        self.assertTrue(out["m1"]["acknowledged"] and out["r1"]["acknowledged"])
        self.assertEqual(self.zoho.orders["m1"]["current_sub_status"], "cs_acknowl")
        # an order with its own cs_acknowl sub-status is acknowledged even when the view has not caught up
        self.zoho.orders["m2"]["current_sub_status"] = self.zoho.orders["m2"]["order_sub_status"] = "cs_acknowl"
        self.reset_caches()
        self.assertTrue(self.acknowledge("m2")["m2"]["already_acknowledged"])


class SourceFilterTests(unittest.TestCase):
    """zoho_client sends branch_ids where Zoho honors it and always drops other branches right after."""

    def run_call(self, fn, payload):
        seen = {}

        def fake_request(method, path, params):
            seen.update(path=path, params=params)
            return payload

        with patch.object(zoho_client, "_request", fake_request):
            return fn(), seen

    ROWS = {"salesorders": [{"salesorder_id": "1", "salesorder_number": "SO26-1", "branch_id": RGF}, {"salesorder_id": "2", "salesorder_number": "SS SO26-2", "branch_id": SSI}, {"salesorder_id": "3", "salesorder_number": "MS-SO-3", "branch_id": MSSI}, {"salesorder_id": "4", "salesorder_number": "X-4", "branch_id": UNKNOWN}], "page_context": {"has_more_page": False}}

    def test_shipment_window_sends_both_branch_ids_in_one_call(self):
        from datetime import date
        out, seen = self.run_call(lambda: zoho_client.fetch_sales_orders_by_shipment_date(date(2026, 10, 10), date(2026, 10, 10)), dict(self.ROWS))
        self.assertEqual(seen["params"]["branch_ids"], f"{RGF},{MSSI}")
        self.assertEqual([r["salesorder_id"] for r in out["salesorders"]], ["1", "3"])
        self.assertEqual(out["raw_count"], 4)

    def test_custom_view_and_filter_by_do_not_send_branch_ids_but_still_filter(self):
        out, seen = self.run_call(lambda: zoho_client.fetch_sales_orders_by_customview("v"), dict(self.ROWS))
        self.assertNotIn("branch_ids", seen["params"])
        self.assertEqual(len(out["salesorders"]), 2)
        out, seen = self.run_call(lambda: zoho_client.fetch_sales_orders(filter_by="Status.Confirmed"), dict(self.ROWS))
        self.assertNotIn("branch_ids", seen["params"])
        self.assertEqual(len(out["salesorders"]), 2)
        _, seen = self.run_call(lambda: zoho_client.fetch_sales_orders(), dict(self.ROWS))
        self.assertEqual(seen["params"]["branch_ids"], f"{RGF},{MSSI}")

    def test_packages_are_filtered_by_branch_and_number(self):
        payload = {"packages": [{"package_id": "p1", "salesorder_number": "SO26-1"}, {"package_id": "p2", "salesorder_number": "SS SO26-2"}, {"package_id": "p3", "salesorder_number": "MS-SO-3"}], "page_context": {}}
        out, seen = self.run_call(lambda: zoho_client.fetch_packages(), payload)
        self.assertEqual(seen["params"]["branch_ids"], f"{RGF},{MSSI}")
        self.assertEqual([p["package_id"] for p in out["packages"]], ["p1", "p3"])


class BranchConfigTests(unittest.TestCase):
    def test_is_allowed(self):
        self.assertTrue(bs.is_allowed({"branch_id": RGF}))
        self.assertTrue(bs.is_allowed({"branch_id": MSSI, "salesorder_number": "anything"}))
        for other in (SSI, RC, RFS, UNKNOWN):
            self.assertFalse(bs.is_allowed({"branch_id": other, "salesorder_number": "SO26-1"}), other)
        self.assertTrue(bs.is_allowed({"salesorder_number": "SO26-1"}))
        self.assertTrue(bs.is_allowed({"salesorder_number": "WM-SO26-00247"}))
        self.assertTrue(bs.is_allowed({"salesorder_number": "MS-SO-02016"}))
        self.assertFalse(bs.is_allowed({"salesorder_number": "SS SO26-16612"}))
        self.assertFalse(bs.is_allowed({}))
        self.assertTrue(bs.is_allowed(SimpleNamespace(raw_json=None, salesorder_number="SO26-5")))
        self.assertFalse(bs.is_allowed(SimpleNamespace(raw_json={"branch_id": SSI}, salesorder_number="SO26-5")))

    def test_branch_of_ignores_number_prefix_for_the_badge(self):
        self.assertIsNone(bs.branch_of({"salesorder_number": "MS-SO-02016"}))  # no branch fields -> no badge
        self.assertEqual(bs.branch_of(SimpleNamespace(raw_json={"branch_id": MSSI}))["code"], "MSSI")
        self.assertIsNone(bs.branch_of(SimpleNamespace(raw_json=None)))

    def test_parse_ids(self):
        self.assertEqual(bs.parse_ids(f" {RGF}, ,{MSSI} "), {RGF, MSSI})
        self.assertEqual(bs.parse_ids(None), set())
        self.assertEqual(bs.parse_ids(object()), set())


WAREHOUSES = {"warehouses": [
    {"warehouse_name": "Mets Cold Storage Services Inc. RGF", "warehouse_available_for_sale_stock": 10},
    {"warehouse_name": "Glacier South RGF", "warehouse_available_for_sale_stock": 20},
    {"warehouse_name": "Chilled - Mets Cold Storage Services Inc. RGF", "warehouse_available_for_sale_stock": 99},
    {"warehouse_name": "METS RGF (Near-Expiry Items / Production Reserves)", "warehouse_available_for_sale_stock": 98},
    {"warehouse_name": "Mets Cold Storage Services Inc. MSSI / SUPERMARKET", "warehouse_available_for_sale_stock": 30},
    {"warehouse_name": "Glacier South MSSI", "warehouse_available_for_sale_stock": -2000},
    {"warehouse_name": "SariSuki Store Inc. Warehouse", "warehouse_available_for_sale_stock": 50},
    {"warehouse_name": "Glacier South MSSI(DEACTIVATED)", "warehouse_available_for_sale_stock": 77},
]}


class StockMappingTests(unittest.TestCase):
    def parse(self, branch_id, enabled):
        with patch.dict(os.environ, {"BRANCH_STOCK_MAPPING": "1" if enabled else ""}):
            return ws._parse_stock(WAREHOUSES, "i1", branch_id)

    def test_flag_off_is_the_unchanged_legacy_logic(self):
        for branch in (None, RGF, MSSI):
            self.assertEqual(self.parse(branch, False), {"mets": 30, "glacier": 20})

    def test_rgf_and_mssi_mapping_when_on_and_negative_is_kept(self):
        self.assertEqual(self.parse(RGF, True), {"mets": 10, "glacier": 20})
        self.assertEqual(self.parse(MSSI, True), {"mets": 30, "glacier": -2000})

    def test_only_rgf_and_mssi_have_rules(self):
        self.assertEqual(set(bs.STOCK_RULES), {RGF, MSSI})
        self.assertEqual(self.parse(SSI, True), self.parse(None, False))  # no SSI row: legacy default, no "other" column


class ExportColumnsTests(unittest.TestCase):
    def test_no_branch_warehouse_column_and_negative_kept(self):
        from openpyxl import load_workbook
        from services.inventory_exports import CONFIRMED_HEADERS, CONFIRMED_WIDTHS, flatten_confirmed_order, make_excel, make_pdf
        self.assertFalse([h for h in CONFIRMED_HEADERS if "Branch" in h])
        order_row = SimpleNamespace(expected_shipment_date=None, salesorder_number="MS-SO-1", customer_name="C", shipping_address=None, raw_json={"branch_id": MSSI, "line_items": [{"item_id": "i1", "name": "Egg", "quantity": 1, "unit": "case"}]})
        rows = flatten_confirmed_order(order_row, lambda item_id: {"mets": -3.0, "glacier": -2000.0}, lambda order: ("", ""))
        self.assertEqual(len(rows[0]), len(CONFIRMED_HEADERS))
        self.assertEqual(len(CONFIRMED_WIDTHS), len(CONFIRMED_HEADERS))
        sheet = load_workbook(make_excel(rows, "cap", "confirmed")).active
        self.assertEqual((sheet.cell(2, 11).value, sheet.cell(2, 12).value), (-3.0, -2000.0))
        self.assertTrue(make_pdf(rows, "cap", "confirmed").getvalue().startswith(b"%PDF"))


if __name__ == "__main__":
    unittest.main()
