"""Delivery method (standard SO field) + Fulfillment type (custom field cf_fulfillment_type): mapping, export and email. No Zoho calls."""
from __future__ import annotations

from datetime import date
from types import SimpleNamespace

from io import BytesIO

from openpyxl import load_workbook

from services import inventory_exports as ex
from services import order_fields


def test_delivery_method_standard_field():
    assert order_fields.delivery_method({"delivery_method": "  Company Truck "}) == "Company Truck"
    assert order_fields.delivery_method({"delivery_method": ""}) is None
    assert order_fields.delivery_method({}) is None
    # a Neon history row has no raw_json, only the delivery_method column
    assert order_fields.delivery_method(SimpleNamespace(delivery_method="Pick-up")) == "Pick-up"


def test_fulfillment_type_every_zoho_shape():
    assert order_fields.fulfillment_type({"cf_fulfillment_type": "Company Delivery"}) == "Company Delivery"
    assert order_fields.fulfillment_type({"custom_field_hash": {"cf_fulfillment_type": "Pick-up by Customer"}}) == "Pick-up by Customer"
    assert order_fields.fulfillment_type({"custom_fields": [{"api_name": "cf_other", "label": "Other", "value": "x"},
                                                           {"api_name": "cf_fulfillment_type", "label": "Fulfillment Type", "value": "3PL"}]}) == "3PL"
    # label match when the api_name differs (e.g. a re-created field)
    assert order_fields.fulfillment_type({"custom_fields": [{"api_name": "cf_ftype_2", "label": "Fulfilment type", "value": "Lalamove"}]}) == "Lalamove"
    assert order_fields.fulfillment_type({"cf_fulfilment_type_new": "Grab"}) == "Grab"
    assert order_fields.fulfillment_type({"custom_fields": [{"label": "Fulfillment Type", "value": ""}]}) is None
    assert order_fields.fulfillment_type({}) is None
    assert order_fields.fulfillment_type(SimpleNamespace(raw_json={"cf_fulfillment_type": "Company Delivery"})) == "Company Delivery"


def test_apply_sets_both_keys_none_when_missing():
    result = order_fields.apply({}, SimpleNamespace(raw_json={}, delivery_method=None))
    assert result == {"delivery_method": None, "fulfillment_type": None}


def _order(**raw):
    return SimpleNamespace(
        id="1", salesorder_number="SO-1", order_type="SO", customer_name="Cust", expected_shipment_date=date(2026, 10, 21),
        shipping_address={}, delivery_method=raw.get("delivery_method"), order_status="confirmed", invoice_status=None,
        payment_status=None, shipment_status=None, total=10, vehicle_id=None,
        raw_json={"line_items": [{"name": "Beef", "sku": "B1", "quantity": 1, "unit": "kg", "warehouse_name": "Mets Cold Storage"}], **raw},
    )


def test_exports_include_both_columns():
    order = _order(delivery_method="Company Truck", cf_fulfillment_type="Company Delivery")
    default = ex.flatten_order(order)[0]
    assert default[ex.HEADERS.index("Delivery Method")] == "Company Truck"
    assert default[ex.HEADERS.index("Fulfillment Type")] == "Company Delivery"
    assert ex.HEADERS[ex.DEFAULT_STATUS_COLUMN] == "Order Status" and ex.HEADERS[ex.DEFAULT_WAREHOUSE_COLUMN] == "Warehouse"
    confirmed = ex.flatten_confirmed_order(order, lambda _i: {}, lambda _o: ("TRK", "Juan"))[0]
    assert confirmed[ex.CONFIRMED_HEADERS.index("Delivery Method")] == "Company Truck"
    assert confirmed[ex.CONFIRMED_HEADERS.index("Fulfillment Type")] == "Company Delivery"
    assert ex.CONFIRMED_HEADERS[ex.CONFIRMED_TRUCK_COLUMN] == "Truck"
    assert len(ex.CONFIRMED_WIDTHS) == len(ex.CONFIRMED_HEADERS) and len(ex.DEFAULT_WIDTHS) == len(ex.HEADERS)
    sheet = load_workbook(BytesIO(ex.make_excel([confirmed], "cap", "confirmed").getvalue())).active
    headers = [sheet.cell(1, c).value for c in range(1, len(ex.CONFIRMED_HEADERS) + 1)]
    assert "Delivery Method" in headers and "Fulfillment Type" in headers
    assert ex.make_pdf([default], "cap").getvalue().startswith(b"%PDF")


def test_assignment_email_has_both_columns():
    from routers import assignment

    order = _order(delivery_method="Company Truck", custom_fields=[{"api_name": "cf_fulfillment_type", "label": "Fulfillment Type", "value": "Company Delivery"}])
    driver = SimpleNamespace(name="Juan", warehouse="METS")
    profile = SimpleNamespace(plate_no="ABC123", capacity_note=None)
    _, html, body = assignment._assignment_email_html([driver], profile, [order], "admin", "now", weight_fn=lambda _o: 1.0, unknown_fn=lambda _o: 0)
    assert ">Delivery method</th>" in body and ">Fulfillment type</th>" in body
    assert "Company Truck" in body and "Company Delivery" in body
    text = assignment._assignment_email_text(profile, [order], weight_fn=lambda _o: 1.0, unknown_fn=lambda _o: 0)
    assert "Delivery method: Company Truck" in text and "Fulfillment type: Company Delivery" in text
