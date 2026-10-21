"""Pickup preview, printable attachments, daily scheduling and durable send outcomes."""
from datetime import datetime, timedelta
from hashlib import sha256
from html import escape
from io import BytesIO
import json
import threading
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from services import app_settings, gmail_sender, pickup_rules, purchase_orders
from models.access_workflow import PickupEmailLog

_send_lock = threading.Lock()
_last_tick = None


def recipients():
    split = lambda key: [v.strip() for v in app_settings.get_str(key).split(',') if v.strip()]
    return {"to": list(dict.fromkeys(split("PICKUP_PROCUREMENT_TO") + split("PICKUP_LOGISTICS_TO"))),
            "cc": list(dict.fromkeys(split("PICKUP_PROCUREMENT_CC") + split("PICKUP_LOGISTICS_CC")))}


def printable(order):
    summary = purchase_orders.summary(order)
    raw = order.raw_json or {}
    address = raw.get('shipping_address') or {}
    contact = raw.get('vendor_contact') or {}
    return {"po": order.salesorder_number, "vendor": order.customer_name,
        "address": ", ".join(str(address[k]) for k in ('address', 'street2', 'city', 'state', 'zip', 'country') if address.get(k)),
        "contact": f"{contact.get('person') or ''} {contact.get('phone') or ''}".strip(),
        "delivery_date": str(order.expected_shipment_date), "destination": order.destination_warehouse or '',
        "items": [{"name": p.get('name'), "qty": p.get('quantity'), "unit": p.get('unit'), "kg": p.get('total_weight_kg')} for p in summary.get('products', [])],
        "vendor_address_missing": raw.get('vendor_address_missing', True)}


def preview():
    day = datetime.now(pickup_rules.PHT).date() + timedelta(days=1)
    records, _ = purchase_orders.get_records()
    scheduled, pahabol = [], []
    for record in records.values():
        if not purchase_orders.is_wanted(record) or pickup_rules.classify(record) == 'excluded':
            continue
        pid = str(record['purchaseorder_id'])
        detail = purchase_orders.cached_detail(pid, record) or purchase_orders.fetch_detail(pid, record)
        merged = {**record, **detail}
        state = purchase_orders.pickup_state(merged)
        if state in {'scheduled', 'pahabol'}:
            (scheduled if state == 'scheduled' else pahabol).append(printable(purchase_orders.build_row(record, detail)))
    scheduled.sort(key=lambda r: r['po'])
    pahabol.sort(key=lambda r: r['po'])
    data = {"delivery_date": day.isoformat(), "recipients": recipients(), "scheduled": scheduled, "pahabol": pahabol}
    data['preview_token'] = sha256(json.dumps(data, sort_keys=True, default=str).encode()).hexdigest()
    return data


def documents(scheduled, pahabol):
    from openpyxl import Workbook
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    styles = getSampleStyleSheet()
    story = [Paragraph('Rare Global Food - Pickup list', styles['Title'])]
    book = Workbook()
    book.remove(book.active)
    for title, rows in [('Scheduled pickups', scheduled), ('Not scheduled - needs Logistics approval', pahabol)]:
        story.extend([Spacer(1, 12), Paragraph(title, styles['Heading2'])])
        sheet = book.create_sheet('Scheduled' if title == 'Scheduled pickups' else 'Pahabol - not scheduled')
        headers = ['PO#', 'Vendor', 'Pickup shipping address', 'Contact / phone', 'Delivery date', 'Destination warehouse', 'Item', 'Qty', 'Unit', 'kg']
        sheet.append(headers)
        sheet.freeze_panes = 'A2'
        for row in rows:
            warning = ' - Vendor address missing' if row.get('vendor_address_missing') else ''
            story.append(Paragraph(escape(f"{row['po']} | {row['vendor']}{warning}"), styles['Heading3']))
            story.append(Paragraph(escape(f"Pick up at {row['address'] or 'Vendor address missing'}. Contact: {row['contact'] or 'Missing'}. Deliver {row['delivery_date']} to {row['destination']}"), styles['BodyText']))
            cells = [['Item', 'Qty', 'Unit', 'kg']]
            for item in row['items'] or [{}]:
                values = [row['po'], row['vendor'], row['address'], row['contact'], row['delivery_date'], row['destination'], item.get('name') or '', item.get('qty'), item.get('unit'), item.get('kg')]
                sheet.append(values)
                # Treat all provider strings as text, including formula-like product/vendor names.
                for cell in sheet[sheet.max_row]:
                    if isinstance(cell.value, str):
                        cell.data_type = 's'
                cells.append([Paragraph(escape(str(item.get(k) if item.get(k) is not None else '—')), styles['BodyText']) for k in ('name','qty','unit','kg')])
            table = Table(cells, colWidths=[450, 70, 80, 80], repeatRows=1, hAlign='LEFT')
            table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e5e7eb')), ('GRID',(0,0),(-1,-1),0.4,colors.grey), ('VALIGN',(0,0),(-1,-1),'TOP')]))
            story.append(table)
        from openpyxl.styles import Alignment, Font
        from openpyxl.utils import get_column_letter
        for col, width in enumerate([23,30,55,30,17,30,45,12,12,12],1):
            sheet.column_dimensions[get_column_letter(col)].width = width
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.fitToWidth = 1
        sheet.print_title_rows = '1:1'
    pdf, excel = BytesIO(), BytesIO()
    SimpleDocTemplate(pdf, pagesize=landscape(A4), rightMargin=28, leftMargin=28).build(story)
    book.save(excel)
    return pdf.getvalue(), excel.getvalue()


PICKUP_TABLE_COLUMNS = ("PO#", "Vendor", "Pickup address", "Contact / phone", "Delivery date", "Destination warehouse", "Items", "Total kg")


def _kg(value):
    try:
        return f"{round(float(value), 2):g}"
    except (TypeError, ValueError):
        return None


def _row_kg(row):
    known = [_kg(item.get('kg')) for item in row.get('items') or []]
    total = sum(float(v) for v in known if v is not None)
    missing = sum(v is None for v in known)
    return (f"{round(total, 2):g} kg" + ("+" if missing else "")) if known and len(known) > missing else "—", missing


def email_html(data):
    """Same RareChain HTML layout as the crew assignment email (inline CSS only; Gmail strips <style>)."""
    from services.rarechain_email_template import render_rarechain_email
    cell = "border:1px solid #cfd4da;padding:8px 10px;text-align:left;vertical-align:top;font-size:13px;line-height:1.4;color:#1f2933;"
    head = "border:1px solid #cfd4da;padding:8px 10px;text-align:left;background:#eef1f4;font-size:12px;font-weight:700;color:#1f2933;white-space:nowrap;"
    headers = "".join(f"<th style='{head}'>{escape(h)}</th>" for h in PICKUP_TABLE_COLUMNS)
    grand_total = 0.0
    sections = []
    for title, note, rows in [("Scheduled pickups", "Confirmed for pickup on the delivery date below.", data['scheduled']),
                              ("Not scheduled – needs Logistics approval", "Pahabol: requested after the cut-off. Logistics must approve before pickup.", data['pahabol'])]:
        body_rows = []
        for row in rows:
            kg_text, missing = _row_kg(row)
            grand_total += sum(float(v) for v in (_kg(i.get('kg')) for i in row.get('items') or []) if v is not None)
            address = escape(row.get('address') or '') if row.get('address') and not row.get('vendor_address_missing') else (
                f"{escape(row.get('address') or '')}<br>" if row.get('address') else "") + "<span style='color:#b42318;font-weight:700;'>Vendor address missing</span>"
            items = "<br>".join(escape(f"{item.get('name') or '—'} — {item.get('qty') if item.get('qty') is not None else '—'} {item.get('unit') or ''}".strip()
                                       + (f" / {_kg(item.get('kg'))} kg" if _kg(item.get('kg')) is not None else " / kg unknown")) for item in row.get('items') or []) or "—"
            kg_cell = escape(kg_text) + (f"<br><span style='font-size:11px;color:#6b7280;'>{missing} line(s) without weight</span>" if missing and kg_text != "—" else "")
            contact = escape(row.get('contact') or '') or "<span style='color:#b42318;'>Missing</span>"
            body_rows.append(f"<tr><td style='{cell}white-space:nowrap;font-weight:700;'>{escape(str(row.get('po') or '—'))}</td><td style='{cell}'>{escape(str(row.get('vendor') or '—'))}</td>"
                             f"<td style='{cell}'>{address}</td><td style='{cell}'>{contact}</td>"
                             f"<td style='{cell}white-space:nowrap;'>{escape(str(row.get('delivery_date') or '—'))}</td><td style='{cell}'>{escape(row.get('destination') or '—')}</td>"
                             f"<td style='{cell}'>{items}</td><td style='{cell}white-space:nowrap;'>{kg_cell}</td></tr>")
        table = (f"<div style='overflow-x:auto;-webkit-overflow-scrolling:touch;'><table cellspacing='0' cellpadding='0' border='0' style='border-collapse:collapse;width:100%;min-width:720px;'>"
                 f"<thead><tr>{headers}</tr></thead><tbody>{''.join(body_rows)}</tbody></table></div>") if body_rows else "<p style='margin:0;font-size:13px;color:#77787B;'>None.</p>"
        sections.append(f"<h2 style='margin:24px 0 4px;font-size:16px;color:#0B0B0B;'>{escape(title)} ({len(rows)})</h2>"
                        f"<p style='margin:0 0 10px;font-size:12px;color:#77787B;'>{escape(note)}</p>{table}")
    body_html = ("<p style='margin:0 0 8px;font-size:14px;line-height:1.5;'>Please review the vendor pickups below and reply to confirm. "
                 "The PDF and Excel versions are attached.</p>" + "".join(sections))
    stats = [{"label": "PICKUP DATE", "value": data['delivery_date']}, {"label": "SCHEDULED POs", "value": str(len(data['scheduled']))},
             {"label": "PAHABOL POs", "value": str(len(data['pahabol']))}, {"label": "TOTAL KG", "value": f"{round(grand_total, 2):g}"}]
    return render_rarechain_email("PROCUREMENT / VENDOR PICKUPS", "Pickup list.", f"Vendor pickups for {data['delivery_date']}.",
                                  "https://images.pexels.com/photos/7464230/pexels-photo-7464230.jpeg?auto=compress&amp;cs=tinysrgb&amp;w=1200", stats,
                                  "https://images.pexels.com/photos/6169056/pexels-photo-6169056.jpeg?auto=compress&amp;cs=tinysrgb&amp;w=1800",
                                  "Pickup details for Procurement and Logistics.", body_html)


def send(db, *, preview_token=None, schedule_key=None):
    with _send_lock:
        data = preview()
        if preview_token is not None and preview_token != data['preview_token']:
            raise ValueError('The pickup list or recipients changed. Preview again before sending.')
        count = len(data['scheduled']) + len(data['pahabol'])
        log = PickupEmailLog(delivery_date=datetime.fromisoformat(data['delivery_date']).date(), recipients=data['recipients'], po_count=count,
                             status='sending', schedule_key=schedule_key)
        db.add(log)
        try:
            db.commit()  # durable claim before external delivery; prevents another scheduler sending twice
        except IntegrityError:
            db.rollback()
            return {"status": "already_attempted"}
        if count == 0:
            log.status = 'skipped_empty'
        elif not data['recipients']['to']:
            log.status, log.error = 'failed', 'Pickup email recipients are not configured'
        else:
            try:
                pdf, excel = documents(data['scheduled'], data['pahabol'])
                html = email_html(data)
                gmail_sender.send_email(to=data['recipients']['to'], cc=data['recipients']['cc'], subject=f"Pickup list — {data['delivery_date']}", html=html,
                    attachments=[('Pickup-list.pdf','application/pdf',pdf), ('Pickup-list.xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',excel)], purpose='pickup-list')
                log.status = 'sent'
            except Exception as exc:
                log.status, log.error = 'failed', str(exc)[:1000]
        db.commit()
        return {"status": log.status, "po_count": count, "error": log.error}


def scheduled_tick(now=None):
    global _last_tick
    now = now or datetime.now(pickup_rules.PHT)
    if app_settings.get_int('PICKUP_EMAIL_ENABLED') != 1 or not recipients()['to']:
        return
    if now.strftime('%H:%M') != app_settings.get_str('PICKUP_EMAIL_TIME'):
        return
    key = now.date().isoformat()
    if _last_tick == key:
        return
    from database import SessionLocal
    with SessionLocal() as db:
        if db.execute(select(PickupEmailLog.id).where(PickupEmailLog.schedule_key == key)).first():
            _last_tick = key
            return
        try:
            send(db, schedule_key=key)
        except Exception as exc:
            db.rollback()
            db.add(PickupEmailLog(delivery_date=now.date()+timedelta(days=1), recipients=recipients(), po_count=0, status='failed', error=str(exc)[:1000], schedule_key=key))
            db.commit()
        _last_tick = key
