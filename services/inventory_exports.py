from __future__ import annotations

from html import escape
from io import BytesIO
from reportlab.lib import colors
from reportlab.lib.pagesizes import A3, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from services import order_types as order_type_rules
from services.sales_order_location import address_lines, shipping_city
from services.warehouses import WAREHOUSES, warehouse_key

# Default layout. "Type" (SO/TO/PO) sits right after the order number; "Warehouse" is the last column
# (it drives the Glacier row fill in the Excel export).
HEADERS = ["Expected Shipment Date", "Order#", "Type", "Customer Name", "Item Description", "SKU", "Quantity", "Unit", "City", "Shipping Address", "Fulfillment Type", "Order Status", "Invoiced", "Payment", "Packed", "Shipped", "Amount", "Warehouse"]
DEFAULT_ORDER_COLUMN = 1
DEFAULT_TYPE_COLUMN = 2
DEFAULT_STATUS_COLUMN = 11
DEFAULT_WAREHOUSE_COLUMN = 17
DEFAULT_WIDTHS = [52, 70, 28, 95, 105, 55, 40, 35, 55, 145, 65, 60, 45, 45, 42, 42, 65, 70]
# Confirmed export: every column the Confirmed screen shows, in the layout dispatch uses.
# "Type" (SO/TO/PO) sits right after the order number column (which holds the SO#/TO#/PO#).
CONFIRMED_HEADERS = ["Date", "Item", "SKU", "Quantity", "Unit", "Order Number", "Type", "Customer", "City", "Shipping Address", "Warehouse", "Mets Qty Available for Sale", "Glacier Qty Available for Sale", "Notes", "Fulfillment Type", "Truck", "Driver/Helper"]
CONFIRMED_QTY_COLUMN = 3
CONFIRMED_ORDER_COLUMN = 5
CONFIRMED_TYPE_COLUMN = 6
CONFIRMED_ADDRESS_COLUMN = 9
CONFIRMED_WAREHOUSE_COLUMN = 10
CONFIRMED_STOCK_COLUMNS = (11, 12)
CONFIRMED_NOTES_COLUMN = 13
CONFIRMED_TRUCK_COLUMN = 15
CONFIRMED_WIDTHS = [50, 100, 58, 42, 30, 62, 28, 85, 52, 118, 68, 52, 56, 150, 62, 56, 80]
CYAN = colors.HexColor("#00FFFF")
NEGATIVE_RED = colors.HexColor("#C00000")
RARE_RED = colors.HexColor("#86000B")
BEIGE = colors.HexColor("#FFEBCE")
HUNTER_GREEN = colors.HexColor("#33673B")
BRAND_WHITE = colors.HexColor("#FFFFFF")
BRAND_BLACK = colors.HexColor("#1B2419")
TYPE_COLORS = {"PO": "#1F4E9E", "SO": "#C00000", "WM": "#7C3AED", "TO": "#2E7D32"}
TYPE_RANK = {"SO": 0, "WM": 1, "TO": 2, "PO": 3}
GLACIER_FILL = "DDEBF7"


def _text(value) -> str:
    return "" if value is None else str(value)

def _date(value) -> str:
    return value.strftime("%d %b %Y") if hasattr(value, "strftime") else _text(value)

def _status_color(value) -> colors.Color:
    normalized = _text(value).lower()
    if normalized in {"confirmed", "acknowledged"}: return HUNTER_GREEN
    if normalized == "void": return RARE_RED
    return BRAND_BLACK

def _address(value) -> str:
    return ", ".join(address_lines(value))

def _fulfillment_type(raw: dict) -> str:
    """Zoho stores this as a custom field (cf_fulfillment_type), not a top-level key."""
    hash_value = (raw.get("custom_field_hash") or {}).get("cf_fulfillment_type") if isinstance(raw.get("custom_field_hash"), dict) else None
    labelled = next((field.get("value") for field in raw.get("custom_fields") or [] if isinstance(field, dict) and str(field.get("label") or "").strip().lower() == "fulfillment type"), None)
    return _text(raw.get("fulfillment_type") or raw.get("cf_fulfillment_type") or hash_value or labelled or raw.get("cf_fulfillment_type_formatted") or raw.get("order_fulfillment_type"))


def _line_item_id(item: dict) -> str | None:
    nested = item.get("item") if isinstance(item.get("item"), dict) else {}
    value = item.get("item_id") or item.get("itemid") or nested.get("item_id") or nested.get("id")
    return str(value) if value else None


def confirmed_item_ids(orders) -> list[str]:
    ids: list[str] = []
    for order in orders:
        for item in ((getattr(order, "raw_json", None) or {}).get("line_items") or []):
            if isinstance(item, dict) and _line_item_id(item):
                ids.append(_line_item_id(item))
    return list(dict.fromkeys(ids))


def _line_warehouse(item: dict) -> str | None:
    nested = item.get("warehouse") if isinstance(item.get("warehouse"), dict) else {}
    return item.get("warehouse_name") or item.get("location_name") or nested.get("warehouse_name") or None


def order_type_of(order) -> str:
    """'SO' | 'WM' | 'PO' | 'TO' (services/order_types.py: a WM-SO number is WM even on a row saved as 'SO'); unknown means SO."""
    value = order_type_rules.order_type_of(order)
    return value if value in TYPE_RANK else "SO"


def type_color(order_type) -> str:
    return TYPE_COLORS.get(str(order_type or "SO").upper(), TYPE_COLORS["SO"])


def order_number_markup(number, order_type) -> str:
    """ReportLab Paragraph markup: ONLY the order number text is coloured."""
    return f'<font color="{type_color(order_type)}">{escape(_text(number))}</font>'


def _source_warehouse_name(order) -> str | None:
    raw = getattr(order, "raw_json", None) or {}
    name = raw.get("from_warehouse_name") or raw.get("source_warehouse")
    if name:
        return _text(name)
    key = raw.get("source_warehouse_key") or getattr(order, "source_warehouse_key", None)
    return WAREHOUSES[key]["name"] if key in WAREHOUSES else None


def _destination_warehouse_name(order) -> str | None:
    raw = getattr(order, "raw_json", None) or {}
    name = raw.get("to_warehouse_name") or getattr(order, "destination_warehouse", None)
    if name:
        return _text(name)
    key = raw.get("destination_warehouse_key")
    return WAREHOUSES[key]["name"] if key in WAREHOUSES else None


def order_line_warehouse(order, item: dict) -> str | None:
    """SO / WM line -> its own line warehouse; TO -> every line comes from the source warehouse; PO -> every line goes to the destination
    (receiving) warehouse."""
    kind = order_type_of(order)
    if kind == "PO":
        return _destination_warehouse_name(order) or _line_warehouse(item)
    own = _line_warehouse(item)
    if kind == "TO":
        return own or _source_warehouse_name(order)
    return own


def row_warehouse_key(row, warehouse_column: int) -> str | None:
    return warehouse_key(row[warehouse_column]) if len(row) > warehouse_column else None


def _truck_label(value) -> str:
    return _text(value).strip()


def sort_confirmed_rows(rows: list[list]) -> list[list]:
    """Group by truck (no truck last), then type SO/WM/TO/PO, then number. Stable: line order within an order is kept."""
    def key(row):
        truck = _truck_label(row[CONFIRMED_TRUCK_COLUMN])
        return (truck == "", truck.casefold(), TYPE_RANK.get(_text(row[CONFIRMED_TYPE_COLUMN]).upper(), 9), _text(row[CONFIRMED_ORDER_COLUMN]).casefold())
    return sorted(rows, key=key)


def group_by_truck(rows: list[list]):
    """Yield ("heading", truck) before each truck's rows, then ("row", row). No headings when no row has a truck."""
    ordered = sort_confirmed_rows(rows)
    if not any(_truck_label(row[CONFIRMED_TRUCK_COLUMN]) for row in ordered):
        for row in ordered:
            yield "row", row
        return
    current = object()
    for row in ordered:
        truck = _truck_label(row[CONFIRMED_TRUCK_COLUMN])
        if truck != current:
            current = truck
            yield "heading", truck or "No truck assigned"
        yield "row", row


def flatten_confirmed_order(order, stock_of, truck_driver) -> list[list]:
    """One row per real Zoho line item, in CONFIRMED_HEADERS order. `stock_of(item_id)` returns
    {"mets", "glacier"} for that item (parsed for this order's branch); `truck_driver(order)` returns (truck, driver/helper).
    Works for TO orders too (customer = destination warehouse, number = TO#)."""
    raw = getattr(order, "raw_json", None) or {}
    items = [item for item in (raw.get("line_items") or []) if isinstance(item, dict) and any(item.get(key) not in (None, "") for key in ("name", "item_description", "sku", "quantity", "unit"))]
    address_value = raw.get("shipping_address") or getattr(order, "shipping_address", None)
    city = _text(shipping_city(order))
    notes = _text(raw.get("notes") or raw.get("note") or raw.get("customer_notes")).strip()
    truck, driver = truck_driver(order)
    otype = order_type_of(order)
    rows = []
    for item in items:
        item_id = _line_item_id(item)
        stock = (stock_of(item_id) if item_id else None) or {}
        rows.append([_date(order.expected_shipment_date), item.get("name") or item.get("item_description"), item.get("sku"), item.get("quantity"), item.get("unit"), order.salesorder_number, otype, order.customer_name, city, _address(address_value), order_line_warehouse(order, item), stock.get("mets"), stock.get("glacier"), notes, _fulfillment_type(raw), truck, driver])
    return rows


def flatten_order(order) -> list[list]:
    raw = order.raw_json or {}
    # Export only real Zoho line items. UI separator rows or empty objects must
    # never become blank spreadsheet/PDF rows.
    source_items = raw.get("line_items") or []
    items = [item for item in source_items if isinstance(item, dict) and any(item.get(key) not in (None, "") for key in ("name", "item_description", "sku", "quantity", "unit", "item_total", "amount"))]
    shipping = _address(raw.get("shipping_address") or order.shipping_address)
    city = _text(shipping_city(order))
    invoice_status = _text(getattr(order, "invoice_status", "") or raw.get("invoice_status") or raw.get("invoiced_status")).lower()
    payment_status = _text(getattr(order, "payment_status", "") or raw.get("payment_status") or raw.get("paid_status")).lower()
    shipment_status = _text(getattr(order, "shipment_status", "") or raw.get("shipment_status") or raw.get("shipping_status") or raw.get("status")).lower()
    invoices = raw.get("invoices") if isinstance(raw.get("invoices"), list) else []
    packages = raw.get("packages") if isinstance(raw.get("packages"), list) else []
    invoiced = "Yes" if any(item.get("is_invoiced") is True or (item.get("quantity") and item.get("quantity_invoiced", 0) >= item.get("quantity")) for item in items) or any(_text(invoice.get("status")).lower() not in {"draft", "void", "cancelled", "canceled"} for invoice in invoices) or "invoiced" in invoice_status or "billed" in invoice_status else "No"
    payment = "Yes" if "paid" in payment_status or "payment received" in payment_status or any(invoice.get("balance") is not None and float(invoice.get("balance") or 0) <= 0 for invoice in invoices) else "No"
    packed = "Yes" if any(item.get("quantity") and item.get("quantity_packed", 0) >= item.get("quantity") for item in items) or any(value in _text(package.get("status") or package.get("detailed_status")).lower() for package in packages for value in ("packed", "fulfilled", "shipped", "delivered")) or any(value in shipment_status for value in ("packed", "shipped", "delivered")) else "No"
    shipped = "Yes" if any(item.get("quantity") and item.get("quantity_shipped", 0) >= item.get("quantity") for item in items) or any(value in _text(package.get("status") or package.get("detailed_status")).lower() for package in packages for value in ("shipped", "fulfilled", "delivered")) or any(value in shipment_status for value in ("shipped", "delivered")) else "No"
    otype = order_type_of(order)
    return [[_date(order.expected_shipment_date), order.salesorder_number, otype, order.customer_name, item.get("name") or item.get("item_description"), item.get("sku"), item.get("quantity"), item.get("unit"), city, shipping, _fulfillment_type(raw), order.order_status, invoiced, payment, packed, shipped, item.get("item_total") or order.total, order_line_warehouse(order, item)] for item in items]

def _number_format(value) -> str:
    whole = isinstance(value, (int, float)) and float(value).is_integer()
    return "#,##0;[Red]-#,##0" if whole else "#,##0.00;[Red]-#,##0.00"


LEGEND_LINES = [
    ("Order number in blue text = Purchase Order (PO)", "1F4E9E", None),
    ("Order number in red text = Sales Order (SO)", "C00000", None),
    ("Order number in purple text = Wet Market order (WM)", "7C3AED", None),
    ("Order number in green text = Transfer Order (TO)", "2E7D32", None),
    ("Glacier rows = light blue row fill", "000000", GLACIER_FILL),
    ("Mets rows = no fill", "000000", None),
]


def _add_legend_sheet(book: Workbook) -> None:
    sheet = book.create_sheet("Legend")
    sheet.cell(1, 1, "Legend").font = Font(bold=True, size=14)
    for offset, (text, color, fill) in enumerate(LEGEND_LINES, 3):
        cell = sheet.cell(offset, 1, text); cell.font = Font(bold=color != "000000", color=color)
        if fill: cell.fill = PatternFill("solid", fgColor=fill)
    sheet.column_dimensions["A"].width = 60


def _style_data_row(sheet, row_number: int, row: list, order_col: int, type_col: int, warehouse_col: int, warehouse_fill: bool) -> None:
    """Only the order-number text is coloured; Glacier rows get a light-blue fill on every cell, Mets rows none."""
    sheet.cell(row_number, order_col + 1).font = Font(bold=True, color=type_color(row[type_col])[1:])
    if warehouse_fill and row_warehouse_key(row, warehouse_col) == "glacier":
        fill = PatternFill("solid", fgColor=GLACIER_FILL)
        for column in range(1, len(row) + 1):
            sheet.cell(row_number, column).fill = fill


def _make_confirmed_excel(rows: list[list], warehouse_fill: bool = True) -> BytesIO:
    book = Workbook(); sheet = book.active; sheet.title = "Sales Orders"
    thin = Side(style="thin", color="000000"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for column, header in enumerate(CONFIRMED_HEADERS, 1):
        cell = sheet.cell(1, column, header); cell.font = Font(bold=True, color="000000"); cell.fill = PatternFill("solid", fgColor="00FFFF"); cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = border
    sheet.row_dimensions[1].height = 48
    row_number = 2
    for kind, payload in group_by_truck(rows):
        if kind == "heading":
            sheet.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=len(CONFIRMED_HEADERS))
            cell = sheet.cell(row_number, 1, f"Truck: {payload}"); cell.font = Font(bold=True, size=12); cell.fill = PatternFill("solid", fgColor="D9D9D9")
            row_number += 1
            continue
        for column, value in enumerate(payload, 1):
            cell = sheet.cell(row_number, column, value); cell.alignment = Alignment(vertical="center", wrap_text=column - 1 in (CONFIRMED_ADDRESS_COLUMN, CONFIRMED_NOTES_COLUMN)); cell.border = border
            if isinstance(value, (int, float)) and column - 1 in (CONFIRMED_QTY_COLUMN, *CONFIRMED_STOCK_COLUMNS): cell.number_format = _number_format(value)
        _style_data_row(sheet, row_number, payload, CONFIRMED_ORDER_COLUMN, CONFIRMED_TYPE_COLUMN, CONFIRMED_WAREHOUSE_COLUMN, warehouse_fill)
        row_number += 1
    sheet.freeze_panes = "A2"
    widths = [13, 30, 14, 11, 8, 15, 7, 28, 14, 48, 20, 16, 16, 60, 18, 14, 22]
    for column, width in enumerate(widths, 1): sheet.column_dimensions[get_column_letter(column)].width = width
    _add_legend_sheet(book)
    output = BytesIO(); book.save(output); output.seek(0); return output


def make_excel(rows: list[list], caption: str, layout: str = "default", warehouse_fill: bool = True) -> BytesIO:
    if layout == "confirmed": return _make_confirmed_excel(rows, warehouse_fill)
    book = Workbook(); sheet = book.active; sheet.title = "Sales Orders"
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS)); sheet.cell(1, 1, "Rare Global Food Trading Corp.").font = Font(bold=True, size=16)
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(HEADERS)); sheet.cell(2, 1, caption)
    for column, header in enumerate(HEADERS, 1):
        cell = sheet.cell(4, column, header); cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="111111")
    for row_number, row in enumerate(rows, 5):
        for column, value in enumerate(row, 1): sheet.cell(row_number, column, value)
        _style_data_row(sheet, row_number, row, DEFAULT_ORDER_COLUMN, DEFAULT_TYPE_COLUMN, DEFAULT_WAREHOUSE_COLUMN, warehouse_fill)
    sheet.freeze_panes = "A5"
    for column in range(1, len(HEADERS) + 1): sheet.column_dimensions[get_column_letter(column)].width = min(34, max(12, max(len(_text(sheet.cell(row, column).value)) for row in range(4, sheet.max_row + 1)) + 2))
    _add_legend_sheet(book)
    output = BytesIO(); book.save(output); output.seek(0); return output


def legend_markup() -> str:
    parts = ", ".join(f'<font color="{TYPE_COLORS[t]}"><b>{label}</b></font>' for t, label in (("SO", "SO (red)"), ("WM", "WM (purple)"), ("PO", "PO (blue)"), ("TO", "TO (green)")))
    return f"Legend - order number colour: {parts}. Glacier rows are filled light blue in the Excel export; Mets rows have no fill."


def make_pdf(rows: list[list], caption: str, layout: str = "default", warehouse_fill: bool = True) -> BytesIO:
    output = BytesIO(); doc = SimpleDocTemplate(output, pagesize=landscape(A3), rightMargin=0.35 * inch, leftMargin=0.35 * inch, topMargin=0.35 * inch, bottomMargin=0.35 * inch)
    styles = getSampleStyleSheet(); title = ParagraphStyle("ExportTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=20, leading=24, textColor=RARE_RED); heading = ParagraphStyle("ExportHeading", parent=styles["Heading2"], fontName="Helvetica", fontSize=13, leading=16, textColor=BRAND_BLACK); caption_style = ParagraphStyle("ExportCaption", parent=styles["Normal"], fontName="Helvetica", fontSize=9, leading=11, textColor=BRAND_BLACK)
    # Paragraph backgrounds must stay unset: ReportLab would paint them only
    # behind the text bounding box. TableStyle owns all cell/row backgrounds.
    cell_style = ParagraphStyle("ExportCell", parent=styles["Normal"], fontName="Helvetica", fontSize=8, leading=9, textColor=BRAND_BLACK)
    header_style = ParagraphStyle("ExportHeader", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=8, leading=9, textColor=BRAND_WHITE)
    story = [Paragraph("<b>Rare Global Food Trading Corp.</b>", title), Paragraph("Sales Orders - Load Planning", heading), Paragraph(escape(caption), caption_style), Spacer(1, 0.05 * inch), Paragraph(legend_markup(), caption_style), Spacer(1, 0.15 * inch)]
    # A3 landscape gives the complete operational export room for all columns.
    confirmed = layout == "confirmed"
    widths = CONFIRMED_WIDTHS if confirmed else DEFAULT_WIDTHS
    headers = CONFIRMED_HEADERS if confirmed else HEADERS
    order_col = CONFIRMED_ORDER_COLUMN if confirmed else DEFAULT_ORDER_COLUMN
    type_col = CONFIRMED_TYPE_COLUMN if confirmed else DEFAULT_TYPE_COLUMN
    confirmed_header_style = header_style.clone("ConfirmedHeader", textColor=BRAND_BLACK, alignment=1)
    negative_style = cell_style.clone("NegativeStock", textColor=NEGATIVE_RED, fontName="Helvetica-Bold")
    data = [[Paragraph(escape(header), confirmed_header_style if confirmed else header_style) for header in headers]]
    heading_rows: list[int] = []
    glacier_rows: list[int] = []
    warehouse_col = CONFIRMED_WAREHOUSE_COLUMN if confirmed else DEFAULT_WAREHOUSE_COLUMN
    items = list(group_by_truck(rows)) if confirmed else [("row", row) for row in rows]
    for kind, row in items:
        if kind == "heading":
            heading_rows.append(len(data))
            data.append([Paragraph(f"<b>Truck: {escape(_text(row))}</b>", cell_style)] + [""] * (len(headers) - 1))
            continue
        cells = []
        if warehouse_fill and row_warehouse_key(row, warehouse_col) == "glacier":
            glacier_rows.append(len(data))  # same light-blue row fill as the Excel export
        for index, value in enumerate(row):
            style = cell_style
            if index == order_col:
                cells.append(Paragraph(order_number_markup(value, row[type_col]), cell_style))
                continue
            if confirmed and index in CONFIRMED_STOCK_COLUMNS and isinstance(value, (int, float)):
                text = f"{value:,.0f}" if float(value).is_integer() else f"{value:,.2f}"
                cells.append(Paragraph(escape(text), negative_style if value < 0 else style))
                continue
            if not confirmed and index == DEFAULT_STATUS_COLUMN:
                style = cell_style.clone(f"status-{_text(value)}")
                style.textColor = _status_color(value)
            cells.append(Paragraph(escape(_text(value)), style))
        data.append(cells)
    table = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    commands = [("BACKGROUND", (0, 0), (-1, 0), CYAN if confirmed else RARE_RED), ("TEXTCOLOR", (0, 0), (-1, 0), BRAND_BLACK if confirmed else BRAND_WHITE), ("TEXTCOLOR", (0, 1), (-1, -1), BRAND_BLACK), ("BACKGROUND", (0, 1), (-1, -1), BRAND_WHITE), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4), ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4), ("GRID", (0, 0), (-1, -1), 0.25, BRAND_BLACK if confirmed else BEIGE)]
    if not confirmed:
        commands.append(("ROWBACKGROUNDS", (0, 1), (-1, -1), [BRAND_WHITE, BEIGE]))
    for glacier_row in glacier_rows:
        commands.append(("BACKGROUND", (0, glacier_row), (-1, glacier_row), colors.HexColor("#" + GLACIER_FILL)))
    for heading_row in heading_rows:
        commands += [("SPAN", (0, heading_row), (-1, heading_row)), ("BACKGROUND", (0, heading_row), (-1, heading_row), colors.HexColor("#D9D9D9"))]
    table.setStyle(TableStyle(commands))
    story.append(table)

    def page_footer(canvas, document):
        canvas.saveState()
        canvas.setStrokeColor(BEIGE)
        canvas.setLineWidth(0.6)
        canvas.line(document.leftMargin, 0.31 * inch, landscape(A3)[0] - document.rightMargin, 0.31 * inch)
        canvas.setFillColor(BRAND_BLACK)
        canvas.setFont("Helvetica", 7)
        canvas.drawString(document.leftMargin, 0.18 * inch, "Rare Global Food Trading Corp. - Sales Orders - Load Planning")
        canvas.drawRightString(landscape(A3)[0] - document.rightMargin, 0.18 * inch, f"Page {canvas.getPageNumber()}")
        canvas.restoreState()

    doc.build(story, onFirstPage=page_footer, onLaterPages=page_footer); output.seek(0); return output
